"""Importación del inventario desde Excel (.xlsx) o CSV.

Tres pasos, usados por la página "Importar inventario" del panel y por el comando
`python manage.py importar_catalogo`:

1. leer_archivo(nombre, datos)  -> filas crudas (dicts de texto) + errores de formato
2. validar_filas(filas)         -> filas limpias (con números ya convertidos) + errores por fila
3. importar_filas(filas, ...)   -> guarda en la base de datos (o solo simula)

Los pasos 1 y 2 NO dependen de Django, así que se pueden probar sin base de datos.
"""
import csv
import io
import unicodedata
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation

COLUMNAS = [
    "categoria", "marca", "modelo", "producto", "color", "material", "codigo_barras",
    "precio", "costo", "stock_minimo", "zona", "nivel", "caja", "cantidad",
]
OBLIGATORIAS = ["categoria", "marca", "modelo", "producto"]
MAX_FILAS = 5000
MAX_BYTES = 5 * 1024 * 1024

# Encabezados alternativos que la gente suele escribir en Excel
ALIAS = {
    "codigo": "codigo_barras",
    "codigo_de_barras": "codigo_barras",
    "sku": "codigo_barras",
    "modelo_de_telefono": "modelo",
    "modelo_telefono": "modelo",
    "nombre": "producto",
    "nombre_producto": "producto",
    "precio_venta": "precio",
    "precio_de_venta": "precio",
    "costo_proveedor": "costo",
    "costo_de_proveedor": "costo",
    "minimo": "stock_minimo",
    "stock_min": "stock_minimo",
    "piezas": "cantidad",
    "stock": "cantidad",
    "existencia": "cantidad",
    "mueble": "zona",
    "ubicacion": "zona",
}


class ErrorArchivo(Exception):
    """El archivo no se puede leer (formato, tamaño o encabezados)."""


def normalizar_encabezado(texto):
    """'Código de barras ' -> 'codigo_barras' (sin acentos, minúsculas, guiones bajos)."""
    texto = unicodedata.normalize("NFKD", str(texto or "")).encode("ascii", "ignore").decode()
    texto = "_".join(texto.strip().lower().replace("-", " ").split())
    return ALIAS.get(texto, texto)


def _celda_a_texto(valor):
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))  # 7501234567890.0 -> "7501234567890"
    return str(valor).strip()


def _filas_desde_tabla(tabla):
    """tabla: lista de filas (listas de valores). La primera fila son los encabezados."""
    if not tabla:
        raise ErrorArchivo("El archivo está vacío.")
    encabezados = [normalizar_encabezado(h) for h in tabla[0]]
    faltan = [c for c in OBLIGATORIAS if c not in encabezados]
    if faltan:
        raise ErrorArchivo(
            "Faltan columnas obligatorias: " + ", ".join(faltan)
            + ". Descarga la plantilla para ver los nombres esperados."
        )
    filas = []
    for numero, valores in enumerate(tabla[1:], start=2):
        textos = [_celda_a_texto(v) for v in valores]
        if not any(textos):
            continue  # fila vacía
        fila = {c: "" for c in COLUMNAS}
        for encabezado, valor in zip(encabezados, textos):
            if encabezado in fila:
                fila[encabezado] = valor
        fila["_fila"] = numero
        filas.append(fila)
    if not filas:
        raise ErrorArchivo("El archivo no tiene productos debajo de los encabezados.")
    if len(filas) > MAX_FILAS:
        raise ErrorArchivo(f"El archivo tiene {len(filas)} filas; el máximo es {MAX_FILAS}.")
    return filas


def leer_archivo(nombre, datos):
    """Lee un .xlsx o .csv (bytes) y devuelve la lista de filas crudas."""
    if len(datos) > MAX_BYTES:
        raise ErrorArchivo("El archivo pesa más de 5 MB.")
    nombre = (nombre or "").lower()
    if nombre.endswith(".xlsx"):
        from openpyxl import load_workbook

        try:
            libro = load_workbook(io.BytesIO(datos), read_only=True, data_only=True)
        except Exception as e:  # archivo dañado o no es Excel
            raise ErrorArchivo("No se pudo abrir el Excel. ¿Está dañado o protegido con contraseña?") from e
        hoja = libro["Inventario"] if "Inventario" in libro.sheetnames else libro.worksheets[0]
        tabla = [list(f) for f in hoja.iter_rows(values_only=True)]
        libro.close()
        return _filas_desde_tabla(tabla)
    if nombre.endswith(".csv"):
        try:
            texto = datos.decode("utf-8-sig")
        except UnicodeDecodeError:
            texto = datos.decode("latin-1")  # CSV guardado como "CSV (delimitado por comas)" en Excel viejo
        try:
            dialecto = csv.Sniffer().sniff(texto[:4096], delimiters=",;")
        except csv.Error:
            dialecto = csv.excel
        return _filas_desde_tabla(list(csv.reader(io.StringIO(texto), dialecto)))
    raise ErrorArchivo("Sube un archivo .xlsx (Excel) o .csv.")


def _decimal(valor):
    limpio = (valor or "0").replace("$", "").replace(",", "").strip() or "0"
    numero = Decimal(limpio)
    if numero < 0:
        raise InvalidOperation
    return numero


def _entero(valor):
    numero = _decimal(valor)
    if numero != numero.to_integral_value():
        raise InvalidOperation
    return int(numero)


def validar_filas(filas):
    """Revisa TODAS las filas y devuelve (filas_limpias, errores).

    errores: lista de (numero_de_fila, mensaje). Si hay errores, no se debe importar nada.
    """
    errores, limpias = [], []
    codigos, variantes = {}, {}
    for f in filas:
        n = f["_fila"]
        errores_fila = [f"falta '{c}'" for c in OBLIGATORIAS if not f[c]]
        numeros = {}
        for campo, conversion in (("precio", _decimal), ("costo", _decimal),
                                  ("stock_minimo", _entero), ("cantidad", _entero)):
            try:
                numeros[campo] = conversion(f[campo])
            except (InvalidOperation, ValueError):
                tipo = "un número entero" if conversion is _entero else "un número"
                errores_fila.append(f"'{campo}' debe ser {tipo} positivo (dice \"{f[campo]}\")")
        if numeros.get("cantidad", 0) > 0 and not f["zona"]:
            errores_fila.append("tiene cantidad pero no 'zona' (¿dónde se guardan las piezas?)")

        codigo = f["codigo_barras"]
        if codigo:
            if codigo.lower() in codigos:
                errores_fila.append(f"el código {codigo} ya aparece en la fila {codigos[codigo.lower()]}")
            else:
                codigos[codigo.lower()] = n
        clave = tuple(f[c].lower() for c in ("categoria", "marca", "modelo", "producto", "color", "material"))
        if all(clave[:4]):
            if clave in variantes:
                errores_fila.append(f"es el mismo producto, color y material que la fila {variantes[clave]}")
            else:
                variantes[clave] = n

        if errores_fila:
            errores.extend((n, e) for e in errores_fila)
        else:
            limpias.append({**f, **numeros})
    return limpias, errores


@dataclass
class Resultado:
    variantes_nuevas: int = 0
    variantes_existentes: int = 0
    ubicaciones_nuevas: int = 0
    piezas: int = 0
    errores: list = field(default_factory=list)  # (fila, mensaje)
    avisos: list = field(default_factory=list)   # (fila, mensaje)
    guardado: bool = False

    @property
    def total_filas(self):
        return self.variantes_nuevas + self.variantes_existentes


class _Simulacion(Exception):
    """Se lanza para deshacer la transacción cuando solo se está revisando."""


def importar_filas(filas, *, usuario, guardar):
    """Importa filas crudas. Con guardar=False hace todo y luego lo deshace (vista previa).

    Si hay cualquier error, no se guarda nada.
    """
    from django.db import transaction

    from catalogo.models import Categoria, Marca, ModeloTelefono, Producto, Variante

    from .models import Movimiento, Ubicacion
    from .services import registrar_movimiento

    resultado = Resultado()
    limpias, resultado.errores = validar_filas(filas)
    if resultado.errores:
        return resultado

    try:
        with transaction.atomic():
            for f in limpias:
                categoria, _ = Categoria.objects.get_or_create(nombre=f["categoria"])
                marca, _ = Marca.objects.get_or_create(nombre=f["marca"])
                modelo, _ = ModeloTelefono.objects.get_or_create(marca=marca, nombre=f["modelo"])
                producto, _ = Producto.objects.get_or_create(
                    categoria=categoria, modelo_telefono=modelo, nombre=f["producto"]
                )
                datos = {"precio": f["precio"], "costo": f["costo"],
                         "stock_minimo": f["stock_minimo"], "creado_por": usuario}
                if f["codigo_barras"]:
                    variante, creada = Variante.objects.get_or_create(
                        codigo_barras=f["codigo_barras"],
                        defaults={"producto": producto, "color": f["color"], "material": f["material"], **datos},
                    )
                else:
                    variante, creada = Variante.objects.get_or_create(
                        producto=producto, color=f["color"], material=f["material"], defaults=datos
                    )
                if creada:
                    resultado.variantes_nuevas += 1
                else:
                    resultado.variantes_existentes += 1

                if f["cantidad"] <= 0:
                    continue
                if not creada:
                    resultado.avisos.append((
                        f["_fila"],
                        f"{variante} ya estaba registrado; no se sumaron sus {f['cantidad']} piezas "
                        "para no duplicar el stock. Registra una entrada si llegaron piezas nuevas.",
                    ))
                    continue
                ubicacion, ub_nueva = Ubicacion.objects.get_or_create(
                    zona=f["zona"], nivel=f["nivel"], caja=f["caja"]
                )
                resultado.ubicaciones_nuevas += int(ub_nueva)
                registrar_movimiento(
                    variante=variante, ubicacion=ubicacion, tipo=Movimiento.Tipo.ENTRADA,
                    motivo=Movimiento.Motivo.ALTA, cantidad=f["cantidad"], usuario=usuario,
                    nota="Importación de inventario",
                )
                resultado.piezas += f["cantidad"]
            if not guardar:
                raise _Simulacion
    except _Simulacion:
        return resultado
    resultado.guardado = True
    return resultado


def plantilla_xlsx():
    """Bytes de un Excel vacío con los encabezados correctos y una fila de ejemplo."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    libro = Workbook()
    hoja = libro.active
    hoja.title = "Inventario"
    hoja.append(COLUMNAS)
    hoja.append(["Fundas", "Apple", "iPhone 15", "Funda uso rudo", "Negro", "Policarbonato",
                 "7501234567890", 250, 90, 3, "Cajonera 1", "A", "Caja Azul", 8])
    for celda in hoja[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="B45309")
    anchos = [14, 12, 18, 26, 12, 16, 16, 9, 9, 13, 12, 7, 12, 10]
    for i, ancho in enumerate(anchos):
        hoja.column_dimensions[chr(65 + i)].width = ancho
    for fila in range(2, 1001):
        hoja.cell(fila, 7).number_format = "@"  # código como texto: evita 7.5E+12
    hoja.freeze_panes = "A2"
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()
